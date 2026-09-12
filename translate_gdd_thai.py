from docx import Document
from docx.oxml.ns import qn

PATH = "Documentation/Clean_and_Learn_GDD_TH.docx"

T = {
"Game Design Document": "เอกสารออกแบบเกม",
"A first person cleaning puzzle game built in Unity": "เกมพัซเซิลทำความสะอาดมุมมองบุคคลที่หนึ่ง พัฒนาด้วย Unity",
"This document presents the current playable design: the core loop, stain and item logic, four mission rooms, two difficulty modes, and player feedback systems.": "เอกสารนี้สรุปการออกแบบเกมฉบับเล่นได้ในปัจจุบัน ครอบคลุม Gameplay Loop ระบบคราบและไอเทม ห้องภารกิจ 4 ห้อง โหมดความยาก 2 รูปแบบ และระบบ Feedback ของผู้เล่น",
"Project role: Game Designer and Unity Developer": "บทบาทในโครงการ: Game Designer และ Unity Developer",
"1 Game Overview": "1 ภาพรวมเกม",
"Clean and Learn is a first person cleaning puzzle game. The player enters a room, locates visible and hidden stains, reads the stain information, chooses an appropriate cleaning item, and clears every target while managing mistakes.": "Clean and Learn เป็นเกมพัซเซิลทำความสะอาดมุมมองบุคคลที่หนึ่ง ผู้เล่นเข้าสู่ห้อง ค้นหาคราบที่มองเห็นและคราบที่ซ่อนอยู่ อ่านข้อมูลคราบ เลือกไอเทมทำความสะอาดที่เหมาะสม และกำจัดคราบทุกจุดภายใต้ข้อจำกัดของความผิดพลาด",
"2 Design Goal": "2 เป้าหมายการออกแบบ",
"The game is designed to make cleaning a decision-making puzzle rather than a simple interaction task. Players are expected to observe each stain, connect its clues to an appropriate item, and consider the consequence of selecting the wrong cleaner.": "เกมออกแบบให้การทำความสะอาดเป็นพัซเซิลที่ต้องตัดสินใจ ไม่ใช่เพียงการกดโต้ตอบ ผู้เล่นต้องสังเกตคราบแต่ละจุด เชื่อมโยงคำใบ้กับไอเทมที่เหมาะสม และพิจารณาผลลัพธ์ของการเลือกน้ำยาผิด",
"Encourage observation before interaction.": "ส่งเสริมให้ผู้เล่นสังเกตก่อนโต้ตอบ",
"Make every cleaning item meaningful through a clear stain-to-tool relationship.": "ทำให้ไอเทมทุกชิ้นมีความหมายผ่านความสัมพันธ์ที่ชัดเจนระหว่างคราบและเครื่องมือ",
"Build difficulty through stain variety, room scale, and limited mistakes.": "เพิ่มความยากด้วยความหลากหลายของคราบ ขนาดของห้อง และจำนวนความผิดพลาดที่จำกัด",
"Keep player feedback visible through the stain counter, item slots, and error status.": "ทำให้ผู้เล่นเห็น Feedback ชัดเจนผ่านตัวนับคราบ ช่องไอเทม และสถานะความผิดพลาด",
"3 Core Gameplay Loop": "3 Core Gameplay Loop",
"4 Player Controls and Interaction": "4 การควบคุมและการโต้ตอบ",
"5 Puzzle Design": "5 การออกแบบพัซเซิล",
"Each puzzle begins with an unknown stain. The player must acquire enough information to choose the correct item rather than using every item at random. A wrong choice is not a hard fail by itself, but it reduces the remaining error allowance and increases pressure in Challenge Mode.": "พัซเซิลทุกจุดเริ่มจากคราบที่ผู้เล่นยังไม่ทราบวิธีแก้ ผู้เล่นต้องเก็บข้อมูลให้เพียงพอก่อนเลือกไอเทม แทนการสุ่มใช้ทุกชิ้น การเลือกผิดไม่ทำให้แพ้ทันที แต่จะลดจำนวนความผิดพลาดที่เหลือและเพิ่มแรงกดดันใน Challenge Mode",
"6 Stain and Item System": "6 ระบบคราบและไอเทม",
"A Cleaning Target stores the stain name, description, required item, discovery state, and clear state. Interaction ranges make stains selectable in the room. The stain visual fades when the correct item is used.": "Cleaning Target เก็บข้อมูลชื่อคราบ คำอธิบาย ไอเทมที่ต้องใช้ สถานะการค้นพบ และสถานะการทำความสะอาด ระยะ Interaction ทำให้ผู้เล่นเลือกคราบในห้องได้ และ Visual ของคราบจะค่อย ๆ หายเมื่อใช้ไอเทมถูกต้อง",
"7 Level Design and Progression": "7 การออกแบบด่านและความก้าวหน้า",
"The four rooms gradually increase the number of active stains and the range of cleaning decisions. Later rooms use wider spaces, more diverse stain types, and more opportunities to make an incorrect selection.": "ห้องทั้ง 4 ห้องเพิ่มจำนวนคราบที่ใช้งานจริงและความหลากหลายของการตัดสินใจทำความสะอาดอย่างต่อเนื่อง ห้องหลัง ๆ มีพื้นที่กว้างขึ้น ประเภทคราบมากขึ้น และมีโอกาสให้เลือกไอเทมผิดมากขึ้น",
"8 Game Modes": "8 โหมดเกม",
"9 UI and Player Feedback": "9 UI และ Player Feedback",
"10 Audio Feedback": "10 Audio Feedback",
"Audio supports the decision loop with distinct feedback for valid and invalid cleaning actions. Correct item use confirms that the player understood the stain logic. Incorrect item use signals the loss of an allowed attempt and reinforces the risk of guessing.": "เสียงสนับสนุนวงจรการตัดสินใจด้วย Feedback ที่แตกต่างกันสำหรับการใช้น้ำยาถูกและผิด การใช้ไอเทมถูกยืนยันว่าผู้เล่นเข้าใจตรรกะของคราบ ส่วนการใช้ผิดจะแจ้งว่าผู้เล่นเสียโอกาสหนึ่งครั้งและย้ำความเสี่ยงของการเดา",
"Correct cleaning use sound": "เสียงใช้ไอเทมทำความสะอาดถูกต้อง",
"Incorrect cleaning use sound": "เสียงใช้ไอเทมทำความสะอาดผิด",
"Menu hover and click sounds": "เสียง Hover และ Click ของเมนู",
"Background music and separate volume controls for music and sound effects": "เพลงประกอบ และการปรับระดับเสียงเพลงกับเอฟเฟกต์แยกกัน",
"11 Technical Implementation": "11 การพัฒนาด้านเทคนิค",
"12 Current Build Summary": "12 สรุป Build ปัจจุบัน",
"The playable build contains four sequential rooms, a stain inspection and cleaning system, item selection through an inventory HUD, mission objectives, normal and challenge configurations, UI guidance, and audio feedback for correct and incorrect item use. The design is structured so that the HUD count reflects real active stains in each mode.": "Build ที่เล่นได้ในปัจจุบันประกอบด้วยห้องเรียงลำดับ 4 ห้อง ระบบตรวจสอบและทำความสะอาดคราบ การเลือกไอเทมผ่าน Inventory HUD วัตถุประสงค์ภารกิจ การตั้งค่า Normal และ Challenge UI แนะนำผู้เล่น และเสียง Feedback สำหรับการใช้ไอเทมถูกหรือผิด โดยตัวเลขใน HUD สอดคล้องกับจำนวนคราบที่ Active จริงในแต่ละโหมด",
"Category": "หมวดหมู่", "Current Design": "การออกแบบปัจจุบัน", "Genre": "แนวเกม", "First person cleaning puzzle": "พัซเซิลทำความสะอาดมุมมองบุคคลที่หนึ่ง", "Engine": "เอนจิน", "Primary platform": "แพลตฟอร์มหลัก", "Windows PC": "Windows PC", "Player objective": "เป้าหมายผู้เล่น", "Clear all stains in a room and complete the mission": "กำจัดคราบทุกจุดในห้องและทำภารกิจให้สำเร็จ", "Progression": "ความก้าวหน้า", "Bathroom, Kitchen, Living Room, Bedroom": "Bathroom, Kitchen, Living Room, Bedroom", "Modes": "โหมด", "Normal Mode and Challenge Mode": "Normal Mode และ Challenge Mode",
"Step": "ขั้นตอน", "Player Action": "การกระทำของผู้เล่น", "Game Feedback": "Feedback ของเกม", "Explore": "สำรวจ", "Move through the room and locate stains or items.": "เดินสำรวจห้องและค้นหาคราบหรือไอเทม", "Environmental placement and item beacons guide attention.": "การจัดวางในสภาพแวดล้อมและ Beacon ของไอเทมช่วยนำสายตา", "Observe": "สังเกต", "Inspect a stain to read its name, description, and clue.": "ตรวจข้อมูลคราบเพื่ออ่านชื่อ คำอธิบาย และคำใบ้", "Information pop up explains the stain and its required item.": "Pop-up ข้อมูลอธิบายคราบและไอเทมที่ต้องใช้", "Interact": "โต้ตอบ", "Pick up an item and select it from the inventory slots.": "เก็บไอเทมและเลือกจากช่อง Inventory", "HUD shows the available item slots and active selection.": "HUD แสดงช่องไอเทมที่มีและช่องที่เลือกอยู่", "Experiment": "ทดลอง", "Use the chosen item on a discovered stain.": "ใช้ไอเทมที่เลือกกับคราบที่ค้นพบ", "Correct and incorrect use have distinct feedback and sound.": "การใช้ถูกและผิดมี Feedback และเสียงที่แตกต่างกัน", "Solve": "แก้ปัญหา", "Clear every required stain without exceeding the error limit.": "กำจัดคราบที่จำเป็นทุกจุดโดยไม่เกินขีดจำกัดความผิดพลาด", "HUD updates stain progress and errors left.": "HUD อัปเดตความคืบหน้าของคราบและ Errors Left", "Progress": "ก้าวหน้า", "Complete the mission and unlock the next room.": "ทำภารกิจสำเร็จและปลดล็อกห้องถัดไป", "Level result and level selection reflect progression.": "ผลลัพธ์ด่านและหน้าเลือกด่านสะท้อนความก้าวหน้าของผู้เล่น",
"Input": "ปุ่ม", "Function": "การทำงาน", "Move the player": "เคลื่อนที่ตัวละคร", "Mouse": "เมาส์", "Look around": "หมุนมุมกล้อง", "Space": "Space", "Jump": "กระโดด", "Inspect stain or item information": "ตรวจข้อมูลคราบหรือไอเทม", "Pick up an item or use the selected item": "เก็บไอเทมหรือใช้ไอเทมที่เลือก", "1 to 5": "1 ถึง 5", "Select an inventory slot": "เลือกช่อง Inventory", "Pause the game or close an open panel": "หยุดเกมหรือปิดหน้าต่างที่เปิดอยู่",
"Puzzle Stage": "ขั้นตอนพัซเซิล", "Purpose": "จุดประสงค์", "Problem": "ปัญหา", "A stain is present but its correct cleaning item is not immediately stated.": "มีคราบอยู่ในห้อง แต่ไม่ได้บอกไอเทมที่ถูกต้องทันที", "Observation": "การสังเกต", "The player inspects the stain description and visual context.": "ผู้เล่นตรวจคำอธิบายคราบและบริบทของตำแหน่ง", "Exploration": "การค้นหา", "The player searches the room for the relevant cleaning item.": "ผู้เล่นค้นหาไอเทมทำความสะอาดที่เกี่ยวข้องในห้อง", "The player applies the selected item to the stain.": "ผู้เล่นใช้ไอเทมที่เลือกกับคราบ", "Consequence": "ผลลัพธ์", "Correct use clears the stain. Incorrect use consumes an allowed mistake.": "ใช้ถูก คราบหาย ใช้ผิด จะเสียโอกาสความผิดพลาดหนึ่งครั้ง", "Solution": "วิธีแก้", "The room is completed when every active stain is cleared.": "ห้องจะสำเร็จเมื่อกำจัดคราบ Active ทุกจุด",
"Stain Type": "ประเภทคราบ", "Typical Room Context": "บริบทของห้อง", "Example Required Item": "ตัวอย่างไอเทมที่ใช้", "Water and splash marks": "คราบน้ำและรอยกระเซ็น", "Bathroom or bedroom furniture": "ห้องน้ำหรือเฟอร์นิเจอร์ในห้องนอน", "Water or a suitable cleaner": "น้ำหรือน้ำยาที่เหมาะสม", "Grease and food residue": "คราบมันและเศษอาหาร", "Kitchen counters and cooking areas": "เคาน์เตอร์และพื้นที่ทำอาหารในครัว", "Grease remover or dish cleaner": "น้ำยาขจัดคราบมันหรือน้ำยาล้างจาน", "Dust, hair, and floor debris": "ฝุ่น เส้นผม และเศษบนพื้น", "Living room and bedroom floors": "พื้นห้องนั่งเล่นและห้องนอน", "Floor cleaner": "น้ำยาถูพื้น", "Makeup and surface marks": "คราบเครื่องสำอางและรอยบนพื้นผิว", "Bedroom furniture or wall surfaces": "เฟอร์นิเจอร์หรือผนังห้องนอน", "Appropriate surface cleaner": "น้ำยาทำความสะอาดพื้นผิวที่เหมาะสม", "Insect traces": "ร่องรอยแมลง", "Kitchen corners and storage areas": "มุมครัวและพื้นที่เก็บของ", "Insect spray": "ยาไล่แมลง",
"Mission": "ภารกิจ", "Room": "ห้อง", "Learning Focus": "สิ่งที่ผู้เล่นเรียนรู้", "Normal": "Normal", "Challenge": "Challenge", "Bathroom": "Bathroom", "Basic stain inspection and first item matching": "การตรวจคราบพื้นฐานและการจับคู่ไอเทมครั้งแรก", "3 stains": "3 คราบ", "4 stains": "4 คราบ", "Kitchen": "Kitchen", "Grease, dish residue, and insect related cleaning choices": "การเลือกทำความสะอาดคราบมัน คราบจาน และร่องรอยแมลง", "5 stains": "5 คราบ", "6 stains": "6 คราบ", "Living Room": "Living Room", "Wider exploration and multiple surface types": "การสำรวจพื้นที่กว้างและพื้นผิวหลายประเภท", "7 stains": "7 คราบ", "Bedroom": "Bedroom", "Mixed stain types and final room knowledge check": "คราบหลายประเภทและการทดสอบความรู้ในห้องสุดท้าย", "8 stains": "8 คราบ",
"Mode": "โหมด", "Player Experience": "ประสบการณ์ผู้เล่น", "Primary Difference": "ความแตกต่างหลัก", "Normal Mode": "Normal Mode", "Introduces the cleaning logic and supports learning through play.": "แนะนำตรรกะการทำความสะอาดและช่วยให้ผู้เล่นเรียนรู้ผ่านการเล่น", "Fewer active stains and more forgiving error management.": "มีคราบ Active น้อยกว่าและจัดการความผิดพลาดได้ผ่อนปรนกว่า", "Challenge Mode": "Challenge Mode", "Tests the same knowledge under greater pressure.": "ทดสอบความรู้เดิมภายใต้แรงกดดันที่มากขึ้น", "One additional active stain per room and stricter error limits.": "เพิ่มคราบ Active หนึ่งจุดต่อห้องและจำกัดความผิดพลาดเข้มงวดขึ้น",
"UI Element": "องค์ประกอบ UI", "Mission HUD": "Mission HUD", "Shows current level, stain files, clean combo, best score, and errors left.": "แสดงด่านปัจจุบัน Stain Files Clean Combo คะแนนสูงสุด และ Errors Left", "Inventory HUD": "Inventory HUD", "Shows five item slots and which item is currently selected.": "แสดงช่องไอเทม 5 ช่องและไอเทมที่เลือกอยู่", "Stain Information Panel": "หน้าต่างข้อมูลคราบ", "Displays the stain name, description, required item clue, and status.": "แสดงชื่อคราบ คำอธิบาย คำใบ้ของไอเทมที่ต้องใช้ และสถานะ", "Mission Brief": "Mission Brief", "Explains the room objective before the player begins.": "อธิบายเป้าหมายของห้องก่อนผู้เล่นเริ่มภารกิจ", "Level Select": "หน้าเลือกด่าน", "Shows completed rooms and allows the player to choose Normal or Challenge play.": "แสดงห้องที่ทำสำเร็จและให้เลือกเล่น Normal หรือ Challenge", "Pause Guide": "คู่มือ Pause", "Restates player controls and the core objective during play.": "ทบทวนปุ่มควบคุมและเป้าหมายหลักระหว่างเล่น",
"System": "ระบบ", "Implementation Purpose": "หน้าที่ในการพัฒนา", "Defines stain data, interaction requirements, clear state, and visual fade.": "กำหนดข้อมูลคราบ เงื่อนไข Interaction สถานะการเคลียร์ และการ Fade ของ Visual", "Manages item pickup, inventory slots, selected item, and use flow.": "จัดการการเก็บไอเทม ช่อง Inventory ไอเทมที่เลือก และลำดับการใช้งาน", "Activates the correct stain set for Normal or Challenge Mode.": "เปิดใช้ชุดคราบที่ถูกต้องสำหรับ Normal หรือ Challenge Mode", "Control mission identity, loading flow, and progression.": "ควบคุมข้อมูลภารกิจ การโหลด และความก้าวหน้าของด่าน", "Build and present menu, mission selection, guides, and briefings.": "สร้างและแสดงเมนู การเลือกภารกิจ คู่มือ และ Mission Brief", "Routes correct and incorrect cleaning feedback sounds.": "เรียกใช้เสียง Feedback ของการทำความสะอาดถูกและผิด",
}

def replace_paragraph(p):
    original = p.text
    if original not in T:
        return
    replacement = T[original]
    if not p.runs:
        p.add_run(replacement)
        return
    p.runs[0].text = replacement
    for run in p.runs[1:]:
        run.text = ""
    r_pr = p.runs[0]._element.get_or_add_rPr()
    r_fonts = r_pr.rFonts
    if r_fonts is None:
        from docx.oxml import OxmlElement
        r_fonts = OxmlElement("w:rFonts")
        r_pr.append(r_fonts)
    r_fonts.set(qn("w:eastAsia"), "Noto Sans Thai")

doc = Document(PATH)
for paragraph in doc.paragraphs:
    replace_paragraph(paragraph)
for table in doc.tables:
    for row in table.rows:
        for cell in row.cells:
            for paragraph in cell.paragraphs:
                replace_paragraph(paragraph)
for section in doc.sections:
    for paragraph in section.header.paragraphs + section.footer.paragraphs:
        replace_paragraph(paragraph)
doc.core_properties.title = "เอกสารออกแบบเกม Clean and Learn"
doc.save(PATH)
