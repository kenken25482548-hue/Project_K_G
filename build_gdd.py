from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from pathlib import Path

OUT = Path("Documentation/Clean_and_Learn_GDD_TH.docx")

NAVY = "123A52"
LIGHT_BLUE = "EAF5FA"
LIGHT_GRAY = "F4F6F7"
GRID = "D9D9D9"

def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)

def set_cell_margins(cell, top=110, start=120, bottom=110, end=120):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for side, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{side}"))
        if node is None:
            node = OxmlElement(f"w:{side}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")

def set_cell_border(cell, color=GRID):
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = qn(f"w:{edge}")
        element = borders.find(tag)
        if element is None:
            element = OxmlElement(f"w:{edge}")
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), "6")
        element.set(qn("w:color"), color)

def set_font(run, size=None, bold=None, color=None):
    run.font.name = "Aptos"
    run._element.rPr.rFonts.set(qn("w:ascii"), "Aptos")
    run._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos")
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Noto Sans Thai")
    if size:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if color:
        run.font.color.rgb = RGBColor.from_string(color)

def add_para(doc, text="", style=None, size=None, bold=None, color=None, align=None, after=6):
    p = doc.add_paragraph(style=style)
    if align is not None:
        p.alignment = align
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.line_spacing = 1.15
    run = p.add_run(text)
    set_font(run, size, bold, color)
    return p

def add_bullet(doc, text):
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.line_spacing = 1.12
    set_font(p.add_run(text), 10.5)
    return p

def add_table(doc, headers, rows, widths=None):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    hdr = table.rows[0].cells
    for i, label in enumerate(headers):
        hdr[i].text = ""
        p = hdr[i].paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        set_font(p.add_run(label), 9.5, True, "FFFFFF")
        set_cell_shading(hdr[i], NAVY)
        set_cell_margins(hdr[i])
        set_cell_border(hdr[i])
        hdr[i].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    for index, row in enumerate(rows):
        cells = table.add_row().cells
        for i, value in enumerate(row):
            cells[i].text = ""
            p = cells[i].paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            set_font(p.add_run(str(value)), 9.2)
            set_cell_shading(cells[i], "FFFFFF" if index % 2 == 0 else LIGHT_BLUE)
            set_cell_margins(cells[i])
            set_cell_border(cells[i])
            cells[i].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    if widths:
        for row in table.rows:
            for i, width in enumerate(widths):
                row.cells[i].width = Inches(width)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    return table

def add_heading(doc, text, level=1):
    p = doc.add_paragraph(style=f"Heading {level}")
    p.paragraph_format.space_before = Pt(12 if level == 1 else 8)
    p.paragraph_format.space_after = Pt(5)
    run = p.add_run(text)
    set_font(run, 16 if level == 1 else 12, True, "000000")
    return p

def main():
    OUT.parent.mkdir(parents=True, exist_ok=True)
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Inches(0.72)
    section.bottom_margin = Inches(0.65)
    section.left_margin = Inches(0.78)
    section.right_margin = Inches(0.78)

    normal = doc.styles["Normal"]
    normal.font.name = "Aptos"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Noto Sans Thai")
    normal.font.size = Pt(10.5)

    header = section.header.paragraphs[0]
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    set_font(header.add_run("CLEAN AND LEARN  |  GAME DESIGN DOCUMENT"), 8.5, True, "505050")
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_font(footer.add_run("Clean and Learn"), 8.5, False, "707070")

    title = doc.add_paragraph(style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_before = Pt(70)
    title.paragraph_format.space_after = Pt(12)
    set_font(title.add_run("Clean and Learn"), 30, True, "000000")
    sub = doc.add_paragraph()
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sub.paragraph_format.space_after = Pt(24)
    set_font(sub.add_run("Game Design Document"), 15, False, "000000")
    add_para(doc, "A first person cleaning puzzle game built in Unity", size=12, align=WD_ALIGN_PARAGRAPH.CENTER, after=8)
    add_para(doc, "This document presents the current playable design: the core loop, stain and item logic, four mission rooms, two difficulty modes, and player feedback systems.", size=10.5, align=WD_ALIGN_PARAGRAPH.CENTER, after=12)
    add_para(doc, "Project role: Game Designer and Unity Developer", size=10.5, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, after=0)
    doc.add_page_break()

    add_heading(doc, "1 Game Overview")
    add_para(doc, "Clean and Learn is a first person cleaning puzzle game. The player enters a room, locates visible and hidden stains, reads the stain information, chooses an appropriate cleaning item, and clears every target while managing mistakes.")
    add_table(doc, ["Category", "Current Design"], [
        ("Genre", "First person cleaning puzzle"),
        ("Engine", "Unity 6"),
        ("Primary platform", "Windows PC"),
        ("Player objective", "Clear all stains in a room and complete the mission"),
        ("Progression", "Bathroom, Kitchen, Living Room, Bedroom"),
        ("Modes", "Normal Mode and Challenge Mode"),
    ], [1.55, 4.95])

    add_heading(doc, "2 Design Goal")
    add_para(doc, "The game is designed to make cleaning a decision-making puzzle rather than a simple interaction task. Players are expected to observe each stain, connect its clues to an appropriate item, and consider the consequence of selecting the wrong cleaner.")
    add_bullet(doc, "Encourage observation before interaction.")
    add_bullet(doc, "Make every cleaning item meaningful through a clear stain-to-tool relationship.")
    add_bullet(doc, "Build difficulty through stain variety, room scale, and limited mistakes.")
    add_bullet(doc, "Keep player feedback visible through the stain counter, item slots, and error status.")

    add_heading(doc, "3 Core Gameplay Loop")
    add_table(doc, ["Step", "Player Action", "Game Feedback"], [
        ("Explore", "Move through the room and locate stains or items.", "Environmental placement and item beacons guide attention."),
        ("Observe", "Inspect a stain to read its name, description, and clue.", "Information pop up explains the stain and its required item."),
        ("Interact", "Pick up an item and select it from the inventory slots.", "HUD shows the available item slots and active selection."),
        ("Experiment", "Use the chosen item on a discovered stain.", "Correct and incorrect use have distinct feedback and sound."),
        ("Solve", "Clear every required stain without exceeding the error limit.", "HUD updates stain progress and errors left."),
        ("Progress", "Complete the mission and unlock the next room.", "Level result and level selection reflect progression."),
    ], [0.9, 2.95, 2.65])

    add_heading(doc, "4 Player Controls and Interaction")
    add_table(doc, ["Input", "Function"], [
        ("W A S D", "Move the player"),
        ("Mouse", "Look around"),
        ("Space", "Jump"),
        ("E", "Inspect stain or item information"),
        ("F", "Pick up an item or use the selected item"),
        ("1 to 5", "Select an inventory slot"),
        ("Esc", "Pause the game or close an open panel"),
    ], [1.25, 5.25])
    doc.add_page_break()

    add_heading(doc, "5 Puzzle Design")
    add_para(doc, "Each puzzle begins with an unknown stain. The player must acquire enough information to choose the correct item rather than using every item at random. A wrong choice is not a hard fail by itself, but it reduces the remaining error allowance and increases pressure in Challenge Mode.")
    add_table(doc, ["Puzzle Stage", "Purpose"], [
        ("Problem", "A stain is present but its correct cleaning item is not immediately stated."),
        ("Observation", "The player inspects the stain description and visual context."),
        ("Exploration", "The player searches the room for the relevant cleaning item."),
        ("Experiment", "The player applies the selected item to the stain."),
        ("Consequence", "Correct use clears the stain. Incorrect use consumes an allowed mistake."),
        ("Solution", "The room is completed when every active stain is cleared."),
    ], [1.35, 5.15])

    add_heading(doc, "6 Stain and Item System")
    add_para(doc, "A Cleaning Target stores the stain name, description, required item, discovery state, and clear state. Interaction ranges make stains selectable in the room. The stain visual fades when the correct item is used.")
    add_table(doc, ["Stain Type", "Typical Room Context", "Example Required Item"], [
        ("Water and splash marks", "Bathroom or bedroom furniture", "Water or a suitable cleaner"),
        ("Grease and food residue", "Kitchen counters and cooking areas", "Grease remover or dish cleaner"),
        ("Dust, hair, and floor debris", "Living room and bedroom floors", "Floor cleaner"),
        ("Makeup and surface marks", "Bedroom furniture or wall surfaces", "Appropriate surface cleaner"),
        ("Insect traces", "Kitchen corners and storage areas", "Insect spray"),
    ], [1.5, 2.55, 2.45])

    add_heading(doc, "7 Level Design and Progression")
    add_para(doc, "The four rooms gradually increase the number of active stains and the range of cleaning decisions. Later rooms use wider spaces, more diverse stain types, and more opportunities to make an incorrect selection.")
    add_table(doc, ["Mission", "Room", "Learning Focus", "Normal", "Challenge"], [
        ("01", "Bathroom", "Basic stain inspection and first item matching", "3 stains", "4 stains"),
        ("02", "Kitchen", "Grease, dish residue, and insect related cleaning choices", "5 stains", "6 stains"),
        ("03", "Living Room", "Wider exploration and multiple surface types", "6 stains", "7 stains"),
        ("04", "Bedroom", "Mixed stain types and final room knowledge check", "7 stains", "8 stains"),
    ], [0.65, 1.15, 2.85, 0.95, 0.95])

    add_heading(doc, "8 Game Modes")
    add_table(doc, ["Mode", "Player Experience", "Primary Difference"], [
        ("Normal Mode", "Introduces the cleaning logic and supports learning through play.", "Fewer active stains and more forgiving error management."),
        ("Challenge Mode", "Tests the same knowledge under greater pressure.", "One additional active stain per room and stricter error limits."),
    ], [1.35, 2.8, 2.4])
    doc.add_page_break()

    add_heading(doc, "9 UI and Player Feedback")
    add_table(doc, ["UI Element", "Purpose"], [
        ("Mission HUD", "Shows current level, stain files, clean combo, best score, and errors left."),
        ("Inventory HUD", "Shows five item slots and which item is currently selected."),
        ("Stain Information Panel", "Displays the stain name, description, required item clue, and status."),
        ("Mission Brief", "Explains the room objective before the player begins."),
        ("Level Select", "Shows completed rooms and allows the player to choose Normal or Challenge play."),
        ("Pause Guide", "Restates player controls and the core objective during play."),
    ], [2.1, 4.5])

    add_heading(doc, "10 Audio Feedback")
    add_para(doc, "Audio supports the decision loop with distinct feedback for valid and invalid cleaning actions. Correct item use confirms that the player understood the stain logic. Incorrect item use signals the loss of an allowed attempt and reinforces the risk of guessing.")
    add_bullet(doc, "Correct cleaning use sound")
    add_bullet(doc, "Incorrect cleaning use sound")
    add_bullet(doc, "Menu hover and click sounds")
    add_bullet(doc, "Background music and separate volume controls for music and sound effects")

    add_heading(doc, "11 Technical Implementation")
    add_table(doc, ["System", "Implementation Purpose"], [
        ("CleaningTarget", "Defines stain data, interaction requirements, clear state, and visual fade."),
        ("PlayerItemSystem", "Manages item pickup, inventory slots, selected item, and use flow."),
        ("MissionDifficultyController", "Activates the correct stain set for Normal or Challenge Mode."),
        ("MissionLevelCatalog and LevelFlowManager", "Control mission identity, loading flow, and progression."),
        ("MainMenuUI and MissionStoryUI", "Build and present menu, mission selection, guides, and briefings."),
        ("GameSFXManager", "Routes correct and incorrect cleaning feedback sounds."),
    ], [2.5, 4.1])

    add_heading(doc, "12 Current Build Summary")
    add_para(doc, "The playable build contains four sequential rooms, a stain inspection and cleaning system, item selection through an inventory HUD, mission objectives, normal and challenge configurations, UI guidance, and audio feedback for correct and incorrect item use. The design is structured so that the HUD count reflects real active stains in each mode.")

    doc.core_properties.title = "Clean and Learn Game Design Document"
    doc.core_properties.subject = "Game design document"
    doc.core_properties.author = "Clean and Learn Team"
    doc.save(OUT)

if __name__ == "__main__":
    main()
